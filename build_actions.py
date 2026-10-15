# -*- coding: utf-8 -*-
"""Messages that need a human reply, as a prioritised worklist.

Takes the free-text replies (anything that is not one of the standard confirm
or decline strings), drops pure acknowledgements, buckets the rest by what the
sender actually needs, and orders the buckets by how badly it blocks them on
the day.

Output "הודעות לטיפול.xlsx" is PII and gitignored.
"""
import csv, os, re, collections
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from zalul_status import (CONFIRM, DECLINE, CONFIRM_CSV, CONFIRM_XLSX,
                          norm_phone, load_status, load_registrations,
                          confirm_asof)

OUT = 'הודעות לטיפול.xlsx'

# פטפוט שלא דורש תשובה. נבדק מול הטקסט המלא אחרי ניקוי רווחים.
NOISE = re.compile(r'^(תודה[ !.]*|תודה רבה[ !.]*|אוקיי|אוקי|אישור|בסדר|סבבה|כן|'
                   r'מעולה|יופי|\.|,|\?|היי|שלום|בוקר טוב|ערב טוב|[\W_]*)$')

# הסדר כאן הוא סדר העדיפות בפלט, מהחוסם ביותר למטריד פחות.
BUCKETS = [
    # "גלים" לבדו תופס את "בת גלים" שהוא שם מוקד, ולכן רק צירופים מפורשים
    ('1 · מזג אוויר ובטיחות',
     r'תחזית|גשם|סערה|ים גבוה|גלים גבוהים|מזג ?ה?אוויר|בטיחות|עדיין יתקיים'),
    ('2 · לא יודע לאן להגיע',
     r'לאיזה חוף|לאן נרשמתי|איזה חוף|משובץ|שובצ|לאן אנחנו רשומים|לאיפה נרשמתי|'
     r'לאיפה להגיע|לא קיבלתי מידע|לא קיבלתי פרטים|נשלחו פרטים|תזכירו לי|'
     r'קבלת פרטים|המוכרז|מה המקצה'),
    ('3 · שיבוץ למוקד שגוי',
     r'לא נרשמנו|נרשמתי ל|נרשמנו ל|ביקשתי להרשם|להירשם ל|במקום זה|'
     r'הסרתם אותי|עדכנו מקום|לא בת גלים|חוף אחר'),
    ('4 · נקודת מפגש או שעה',
     r'נקודת מפגש|נק.? ציון|וייז|waze|באיזו שעה|מה השעה|מתי מתחיל|שעת |היכן |'
     r'איפה נפגש|כולל שעה'),
    ('5 · ביקש שיחזרו אליו',
     r'תחזרו אלי|להתכתבות|ללא מענה|לא מגיב|לשוחח עם נציג|אשמח לשיחה|'
     r'אשמח שתחזרו|מי איש הקשר|איש קשר|האחראית|לשאול פה שאלה'),
    ('6 · איחור או הגעה חלקית',
     r'באיחור|להקדים|לחלק מהזמן|רק אחר הצהריים|רק אחרי|צריכה לצאת|'
     r'אוכל להגיע בסביבות|נשתדל להגיע'),
    ('7 · שינוי כמות משתתפים',
     r'נגיע רק|אנחנו \d+ אנשים|אנחנו שני|רק \d+ אנשים|נגיע \d+'),
    ('8 · שינוי תאריך',
     r'15\.10|שינינו את השעה|תאריכים נוספים|תאריך נוסף|יהיה עוד תאריך|'
     r'ביום אחר|שישי הקודם|לא בשישי|ליום חמישי|לא שבת|עובדת בשישי'),
    ('9 · קבוצה או בית ספר',
     r'כית|בית ספר|תלמידים|תלמידות|קבוצה|גן |מורה|הסע'),
    ('10 · ציוד וחניה',
     r'ציוד|כפפות|שקיות|מים|כובע|חניה|חנייה|תחבורה|אוטובוס|רכבת|נעלי|'
     r'מאיזה גיל|ילדים|צלילה'),
    ('11 · ביטול בניסוח חופשי',
     r'לא נוכל|לא אוכל|לא נגיע|לא אגיע|לא משתתף|לא נשתתף|לא יכולים|מבטל|לבטל|'
     r'פעם הבאה|בחו"ל|בחול|לצערי|לא מאשר|לא מגיעים|לא נצליח|מצטער|מצטערת'),
]

status, lastmsg, manual = load_status()
reg = {x['phone']: x for x in load_registrations()[0]}
ASOF = confirm_asof()


def rows_with_identity():
    if os.path.exists(CONFIRM_CSV):
        with open(CONFIRM_CSV, encoding='utf-16', newline='') as f:
            for r in list(csv.reader(f, delimiter='\t'))[1:]:
                if r and r[0].strip():
                    yield {'phone': r[0], 'email': r[1] if len(r) > 1 else '',
                           'name': r[2] if len(r) > 2 else '',
                           'msg': (r[4] if len(r) > 4 else '').strip(),
                           'when': r[5] if len(r) > 5 else ''}
    else:
        wb = openpyxl.load_workbook(CONFIRM_XLSX, read_only=True, data_only=True)
        for r in list(wb[wb.sheetnames[0]].iter_rows(values_only=True))[1:]:
            if r and r[0]:
                yield {'phone': r[0], 'email': r[1] or '', 'name': r[2] or '',
                       'msg': str(r[4] or '').strip(), 'when': str(r[5] or '')}


free = [r for r in rows_with_identity()
        if r['msg'] not in CONFIRM | DECLINE
        and not NOISE.match(' '.join(r['msg'].split()))]

STATUS_HE = {'מאשר': 'אישר/ה', 'מסרב': 'ביטל/ה', 'אחר': 'רק הודעה חופשית'}
items, unbucketed = [], 0
for r in free:
    ph = norm_phone(r['phone'])
    x = reg.get(ph)
    one = ' '.join(r['msg'].split())
    label = next((lbl for lbl, pat in BUCKETS if re.search(pat, one)), '12 · אחר')
    if label == '12 · אחר':
        unbucketed += 1
    items.append({'bucket': label, 'when': str(r['when'])[:16], 'name': r['name'],
                  'phone': ph, 'email': r['email'],
                  'site': x['site'] if x else '(אין רישום תואם)',
                  'people': x['people'] if x else '',
                  'status': STATUS_HE.get(status.get(ph), 'לא השיב/ה'),
                  'msg': one})
ORDER = {lbl: n for n, (lbl, _) in enumerate(BUCKETS)}
items.sort(key=lambda i: (ORDER.get(i['bucket'], 99), i['when']))

# ---------- workbook ----------
HEAD_FILL = PatternFill('solid', fgColor='0871B9')
HEAD_FONT = Font(bold=True, color='FFFFFF', size=11, name='Assistant')
TITLE = Font(bold=True, size=15, color='055A95', name='Assistant')
SUB = Font(size=10, color='5B6B78', name='Assistant')
BODY = Font(size=11, name='Assistant')
BOLD = Font(bold=True, size=11, name='Assistant')
GREY = Font(size=11, name='Assistant', color='999999', italic=True)
BUCKET_FILL = PatternFill('solid', fgColor='EAF4FB')
URGENT_FILL = PatternFill('solid', fgColor='FFF0EE')
THIN = Border(bottom=Side('thin', color='DDE6ED'))
COLS = [('קטגוריה', 24), ('מתי', 17), ('שם', 22), ('טלפון', 14),
        ('מוקד', 30), ('משת׳', 7), ('סטטוס', 16), ('ההודעה', 80), ('מייל', 28)]

wb = Workbook()
ws = wb.active
ws.title = 'לטיפול'
ws.sheet_view.rightToLeft = True
ws['A1'] = 'הודעות הדורשות טיפול · יום הניקיון הגדול 16/10/2026'
ws['A1'].font = TITLE
ws['A2'] = ('%d הודעות · נכון לייצוא %s · ממוין לפי דחיפות, 1 = חוסם ביותר'
            % (len(items), ASOF))
ws['A2'].font = SUB

for c, (h, w) in enumerate(COLS, start=1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    cell.alignment = Alignment(horizontal='center')
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = 'A5'

row, prev = 5, None
for it in items:
    if it['bucket'] != prev:
        prev = it['bucket']
        b = ws.cell(row=row, column=1, value=it['bucket'])
        b.font = BOLD
        n = sum(1 for j in items if j['bucket'] == it['bucket'])
        ws.cell(row=row, column=2, value='%d הודעות' % n).font = SUB
        for c in range(1, len(COLS) + 1):
            ws.cell(row=row, column=c).fill = BUCKET_FILL
        row += 1
    vals = [it['bucket'], it['when'], it['name'], it['phone'], it['site'],
            it['people'], it['status'], it['msg'], it['email']]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = GREY if it['status'] == 'ביטל/ה' else BODY
        cell.border = THIN
        if c == 4:
            cell.number_format = '@'
            cell.alignment = Alignment(horizontal='left')
        elif c in (6,):
            cell.alignment = Alignment(horizontal='center')
        elif c == 8:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    if ORDER.get(it['bucket'], 99) <= 4:          # חמש הקטגוריות החוסמות
        for c in range(1, len(COLS) + 1):
            if it['status'] != 'ביטל/ה':
                ws.cell(row=row, column=c).fill = URGENT_FILL
    row += 1

ws.auto_filter.ref = 'A4:%s%d' % (get_column_letter(len(COLS)), row - 1)
wb.save(OUT)

print('נכתב: %s  (%d הודעות)' % (OUT, len(items)))
by = collections.Counter(i['bucket'] for i in items)
for k in sorted(by):
    print('   %-26s %3d' % (k, by[k]))
