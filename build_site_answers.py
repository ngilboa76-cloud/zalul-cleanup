# -*- coding: utf-8 -*-
"""Yes / no answers for specific sites, split into three clear groups.

Usage:  python build_site_answers.py "בצת" "חשמל"
Each argument is matched as a substring against the site name, so a partial
name is enough. Falls back to the two sites asked for if none are given.

Output "תשובות לפי מוקד.xlsx" is PII and gitignored.
"""
import sys, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from zalul_status import (load_status, load_registrations, load_site_totals,
                          reg_date, confirm_asof, BLAST)

OUT = 'תשובות לפי מוקד.xlsx'
WANTED = sys.argv[1:] or ['בצת', 'חשמל']

status, lastmsg, manual = load_status()
site_totals = load_site_totals()
regs, _ = load_registrations()
ASOF = confirm_asof()

by_site = collections.defaultdict(list)
for x in regs:
    by_site[x['site']].append(x)

targets = [s for s in by_site if any(w in s for w in WANTED)]
if not targets:
    raise SystemExit('לא נמצא מוקד שמכיל: %s' % ', '.join(WANTED))


def group_of(x):
    st = status.get(x['phone'])
    if st == 'מאשר':
        return 'מגיע'
    if st == 'מסרב':
        return 'לא מגיע'
    return 'לא השיב'


GREEN = PatternFill('solid', fgColor='EAF7F0')
RED = PatternFill('solid', fgColor='FDEEEB')
GREYF = PatternFill('solid', fgColor='F2F5F7')
HEAD_FILL = PatternFill('solid', fgColor='0871B9')
HEAD_FONT = Font(bold=True, color='FFFFFF', size=11, name='Assistant')
TITLE = Font(bold=True, size=15, color='055A95', name='Assistant')
SUB = Font(size=10, color='5B6B78', name='Assistant')
BODY = Font(size=11, name='Assistant')
BOLD = Font(bold=True, size=12, name='Assistant')
GREY = Font(size=11, name='Assistant', color='888888')
THIN = Border(bottom=Side('thin', color='DDE6ED'))
COLS = [('שם', 28), ('טלפון', 15), ('משתתפים', 11), ('הערה אחרונה', 46)]
GROUPS = [('מגיע', 'מגיעים', GREEN), ('לא מגיע', 'לא מגיעים', RED),
          ('לא השיב', 'לא השיבו', GREYF)]
BAD = set(r':\/?*[]')

wb = Workbook()
wb.remove(wb.active)
summary = []

for site in sorted(targets, key=lambda s: -sum(x['people'] for x in by_site[s])):
    rows = by_site[site]
    buckets = collections.defaultdict(list)
    for x in rows:
        buckets[group_of(x)].append(x)

    tab = ''.join(' ' if c in BAD else c for c in site).strip()[:31]
    ws = wb.create_sheet(tab)
    ws.sheet_view.rightToLeft = True
    ws['A1'] = site
    ws['A1'].font = TITLE
    tot = sum(x['people'] for x in rows)
    ov = site_totals.get(site)
    counts = {k: sum(x['people'] for x in buckets[k]) for k, _, _ in GROUPS}
    ws['A2'] = ('%d נרשמים · %d משתתפים · מגיעים %d · לא מגיעים %d · לא השיבו %d%s'
                % (len(rows), tot, counts['מגיע'], counts['לא מגיע'], counts['לא השיב'],
                   ' · צפי מהשטח %d' % ov[0] if ov else ''))
    ws['A2'].font = SUB
    ws['A3'] = 'נכון לייצוא %s' % ASOF
    ws['A3'].font = SUB

    for c, (_, w) in enumerate(COLS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    row = 5
    for key, label, fill in GROUPS:
        people = sorted(buckets[key], key=lambda x: (-x['people'], x['name']))
        hdr = ws.cell(row=row, column=1,
                      value='%s  (%d נרשמים · %d משתתפים)'
                            % (label, len(people), counts[key]))
        hdr.font = BOLD
        for c in range(1, len(COLS) + 1):
            ws.cell(row=row, column=c).fill = fill
        row += 1
        for c, (h, _) in enumerate(COLS, start=1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.fill, cell.font = HEAD_FILL, HEAD_FONT
            cell.alignment = Alignment(horizontal='center')
        row += 1
        if not people:
            ws.cell(row=row, column=1, value='(אין)').font = GREY
            row += 2
            continue
        for x in people:
            note = lastmsg.get(x['phone'], '') or ''
            if x['phone'] in manual:
                note = 'אושר בטלפון · ' + note
            elif status.get(x['phone']) is None:
                d = reg_date(x['date'])
                note = 'נרשם/ה אחרי שהמסרון יצא' if d and d >= BLAST else ''
            for c, v in enumerate([x['name'], x['phone'], x['people'],
                                   ' '.join(str(note).split())[:90]], start=1):
                cell = ws.cell(row=row, column=c, value=v)
                cell.font = BODY
                cell.border = THIN
                cell.fill = fill
                if c == 2:
                    cell.number_format = '@'
                    cell.alignment = Alignment(horizontal='left')
                elif c == 3:
                    cell.alignment = Alignment(horizontal='center')
            row += 1
        row += 1

    summary.append((site, len(rows), tot, counts['מגיע'], counts['לא מגיע'],
                    counts['לא השיב']))

wb.save(OUT)
print('נכתב: %s' % OUT)
print('%-26s %6s %6s %7s %8s %8s' % ('מוקד', 'טפסים', 'משת׳', 'מגיעים', 'לא מגיעים', 'לא השיבו'))
print('-' * 68)
for s, n, t, y, no, q in summary:
    print('%-26s %6d %6d %7d %8d %8d' % (s[:26], n, t, y, no, q))
