# -*- coding: utf-8 -*-
"""Shared loading for the attendance pipeline.

build_attendance.py and build_callsheet.py both need the same phone
normalisation, headcount parsing and confirmation status, so it lives here
once instead of drifting between two copies.
"""
import csv, os, re, datetime, collections
import openpyxl

CONFIRM_CSV = '_confirm.csv'        # UTF-16, tab separated (newer export)
CONFIRM_XLSX = '_confirm.xlsx'      # older export format
REGS = '_regs_full.xlsx'
OVERRIDES = 'manual_confirms.csv'       # אישורים שהגיעו בטלפון ולא דרך המסרון
SITE_TOTALS = 'manual_site_totals.csv'  # צפי מגיעים שנקבע ברמת המוקד

CONFIRM = {'מאשר/ת הגעה', 'מאשרת הגעה', 'מאשר הגעה'}
# המערכת שולחת שתי תשובות סירוב תקניות, לא אחת. "לא אגיע" הופיע 93 פעמים
# ונספר בטעות כפנייה חופשית עד שנמצא. כל ערך חדש כאן חייב להיות מחרוזת
# מדויקת מהייצוא, לא ניחוש, אחרת נסווג בטעות הודעות אמיתיות.
DECLINE = {'לא אוכל להגיע', 'לא אגיע'}
TEST_MARKERS = ('בדיקה', 'eddr666')

# המסרון יצא ב-06/10. מי שנרשם אחרי כן פשוט לא היה ברשימת התפוצה.
BLAST = datetime.datetime(2026, 10, 6)


def norm_phone(s):
    """Local Israeli form: digits only, 972 prefix stripped, leading 0 restored."""
    d = re.sub(r'\D', '', str(s or ''))
    if d.startswith('972'):
        d = '0' + d[3:]
    elif d and not d.startswith('0'):
        d = '0' + d
    return d


def people_of(v):
    """Headcount cell. Sheets returns these as floats, so '30.0' must NOT be
    digit-stripped to '300'."""
    if isinstance(v, bool) or v is None or v == '':
        return 1
    if isinstance(v, (int, float)):
        return max(1, int(round(v)))
    return int(re.sub(r'\D', '', str(v)) or 1)


def reg_date(v):
    m = re.match(r'(\d{1,2})/(\d{1,2})/(\d{4})', str(v or ''))
    return datetime.datetime(int(m[3]), int(m[2]), int(m[1])) if m else None


def _confirm_rows():
    """The platform exports either .xlsx or a UTF-16 TAB-separated .csv whose
    message field may contain newlines. Prefer the .csv, it is the newer format."""
    if os.path.exists(CONFIRM_CSV):
        with open(CONFIRM_CSV, encoding='utf-16', newline='') as f:
            for r in list(csv.reader(f, delimiter='\t'))[1:]:
                if r and r[0].strip():
                    yield r[0], (r[5] if len(r) > 5 else ''), (r[4] if len(r) > 4 else '')
    else:
        wb = openpyxl.load_workbook(CONFIRM_XLSX, read_only=True, data_only=True)
        for r in list(wb[wb.sheetnames[0]].iter_rows(values_only=True))[1:]:
            if r and r[0]:
                yield r[0], r[5], (r[4] or '')


def confirm_asof():
    """Date of the newest message in the confirmation export, as DD/MM. This is
    the real as-of date of the confirmation figures, so the dashboard reads it
    from the data instead of carrying a hardcoded date that goes stale."""
    newest = ''
    for _, when, _ in _confirm_rows():
        s = str(when or '')
        if s > newest:
            newest = s
    m = re.match(r'(\d{4})-(\d{2})-(\d{2})', newest)
    return '%s/%s' % (m[3], m[2]) if m else ''


def load_status():
    """phone -> 'מאשר' | 'מסרב' | 'אחר'. Latest decisive message wins; a manual
    override in manual_confirms.csv beats anything from the SMS export, because
    a phone call is newer information than the blast."""
    events = collections.defaultdict(list)
    for phone, when, msg in _confirm_rows():
        msg = str(msg or '').strip()
        kind = 'מאשר' if msg in CONFIRM else 'מסרב' if msg in DECLINE else 'אחר'
        events[norm_phone(phone)].append((str(when), kind, msg))

    status, lastmsg, manual = {}, {}, {}
    for ph, evs in events.items():
        evs.sort(key=lambda e: e[0])
        decisive = [e for e in evs if e[1] != 'אחר']
        status[ph] = decisive[-1][1] if decisive else 'אחר'
        lastmsg[ph] = evs[-1][2]

    if os.path.exists(OVERRIDES):
        with open(OVERRIDES, encoding='utf-8-sig', newline='') as f:
            for row in csv.DictReader(f):
                ph = norm_phone(row.get('טלפון'))
                st = (row.get('סטטוס') or '').strip()
                if not ph or st not in ('מאשר', 'מסרב'):
                    continue
                status[ph] = st
                manual[ph] = (row.get('מקור') or '').strip()
                note = (row.get('הערה') or '').strip()
                if note:
                    lastmsg[ph] = note
    return status, lastmsg, manual


def load_site_totals():
    """site -> (expected_people, note). A figure set from the field for the whole
    site, which can legitimately exceed the registered headcount: groups grow and
    people turn up without registering. Overrides the computed confirmation sum
    for that site only."""
    out = {}
    if not os.path.exists(SITE_TOTALS):
        return out
    with open(SITE_TOTALS, encoding='utf-8-sig', newline='') as f:
        for row in csv.DictReader(f):
            site = (row.get('מוקד') or '').strip()
            try:
                n = int(re.sub(r'\D', '', str(row.get('צפי מגיעים') or '')))
            except ValueError:
                continue
            if site:
                out[site] = (n, (row.get('הערה') or '').strip())
    return out


def load_registrations():
    """One row per (phone, site). When the same person submits the form twice
    for the same site the LAST submission wins: a re-submission is a correction,
    not a duplicate. One registrant cut her group from 10 to 5 that way."""
    wb = openpyxl.load_workbook(REGS, read_only=True, data_only=True)
    keep, skipped = {}, 0
    for r in list(wb['all'].iter_rows(values_only=True))[1:]:
        if not r or not r[0]:
            continue
        name, email = str(r[5] or ''), str(r[6] or '').lower()
        if any(m in name or m in email for m in TEST_MARKERS):
            skipped += 1
            continue
        # כמה הרשמות קבוצתיות הגיעו בלי מוקד כלל. אסור לבלוע אותן בשקט.
        site = (r[4] or r[1] or '').strip() or 'ללא מוקד משויך'
        x = {'date': r[0], 'site': site, 'name': r[5] or '', 'email': r[6] or '',
             'phone': norm_phone(r[7]), 'org': r[8] or '',
             'people': people_of(r[9]), 'type': r[10] or ''}
        keep[(x['phone'], site)] = x
    return list(keep.values()), skipped
