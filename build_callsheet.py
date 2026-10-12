# -*- coding: utf-8 -*-
"""Call sheets for the sites with zero confirmations, one file per site.

Writes a separate .xlsx per site into callsheets_pii/, plus a summary file.
xlsx rather than csv so the leading zero of a phone number survives: Excel
silently eats it when opening a csv.

Status is three-way, not two-way. The SMS went out on 06/10, so anyone who
registered after that was never on the send list. Calling them is a first
ask, not a chase, and the script labels them accordingly.
"""
import os, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from zalul_status import (load_status, load_registrations, load_site_totals,
                          reg_date, BLAST)

OUTDIR = 'callsheets_pii'

status, lastmsg, manual = load_status()
site_totals = load_site_totals()


def status_of(x):
    st = status.get(x['phone'])
    if st == 'מסרב':
        return 'ביטל/ה'
    if st == 'מאשר':
        return 'אישר/ה בטלפון' if x['phone'] in manual else 'אישר/ה'
    if st == 'אחר':
        return 'שלח/ה הודעה חופשית'
    d = reg_date(x['date'])
    if d and d >= BLAST:
        return 'נרשם/ה אחרי המסרון'          # מעולם לא נשאל/ה
    return 'לא השיב/ה כלל'


by_site = collections.defaultdict(lambda: {'yes': 0, 'ppl': 0, 'rows': []})
for x in load_registrations()[0]:
    s = by_site[x['site']]
    s['ppl'] += x['people']
    s['rows'].append(x)
    if status.get(x['phone']) == 'מאשר':
        s['yes'] += x['people']

# מוקד עם צפי מהשטח כבר לא "ללא אישורים", גם אם אף אחד לא השיב למסרון
order = sorted(((k, v) for k, v in by_site.items()
                if v['yes'] == 0 and not site_totals.get(k)),
               key=lambda kv: -kv[1]['ppl'])

# ---- styling ----
HEAD_FILL = PatternFill('solid', fgColor='0871B9')
HEAD_FONT = Font(bold=True, color='FFFFFF', size=11, name='Assistant')
TITLE = Font(bold=True, size=15, color='055A95', name='Assistant')
SUB = Font(size=10, color='5B6B78', name='Assistant')
BODY = Font(size=11, name='Assistant')
GREY = Font(size=11, name='Assistant', color='999999', italic=True)
GROUP_FILL = PatternFill('solid', fgColor='FFF4E0')     # קבוצות, שיחה אחת = הרבה אנשים
NEW_FILL = PatternFill('solid', fgColor='EAF4FB')       # נרשמו אחרי המסרון
THIN = Border(bottom=Side('thin', color='DDE6ED'))
COLS = [('שם', 26), ('טלפון', 14), ('משתתפים', 11), ('סוג', 15),
        ('ארגון', 18), ('מייל', 30), ('תאריך הרשמה', 17), ('סטטוס', 22)]
BAD = r':\/?*[]"<>|'


def safe_name(raw):
    n = ''.join('-' if c in BAD else c for c in raw).strip().rstrip('.')
    return (n[:70] or 'מוקד')


os.makedirs(OUTDIR, exist_ok=True)
for old in os.listdir(OUTDIR):
    if old.endswith('.xlsx'):
        os.remove(os.path.join(OUTDIR, old))

written = []
for rank, (site, s) in enumerate(order, start=1):
    rows = sorted(s['rows'], key=lambda x: -x['people'])
    never = sum(1 for x in rows if status_of(x) == 'לא השיב/ה כלל')
    newreg = sum(1 for x in rows if status_of(x) == 'נרשם/ה אחרי המסרון')
    declined = sum(1 for x in rows if status_of(x) == 'ביטל/ה')

    wbx = Workbook()
    ws = wbx.active
    ws.title = 'רשימת טלפונים'
    ws.sheet_view.rightToLeft = True
    ws['A1'] = site
    ws['A1'].font = TITLE
    ws['A2'] = ('%d משתתפים רשומים · %d הרשמות · אף אחד לא אישר הגעה' % (s['ppl'], len(rows)))
    ws['A2'].font = SUB
    ws['A3'] = ('%d לא השיבו למסרון · %d נרשמו אחרי שהמסרון יצא · %d ביטלו · נכון ל־12/10/2026'
                % (never, newreg, declined))
    ws['A3'].font = SUB

    for c, (h, w) in enumerate(COLS, start=1):
        cell = ws.cell(row=5, column=c, value=h)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = 'A6'

    for i, x in enumerate(rows, start=6):
        st = status_of(x)
        for c, v in enumerate([x['name'], x['phone'], x['people'], x['type'],
                               x['org'], x['email'], x['date'], st], start=1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font = GREY if st == 'ביטל/ה' else BODY
            cell.border = THIN
            if c == 2:
                cell.number_format = '@'                 # keep the leading zero
                cell.alignment = Alignment(horizontal='left')
            elif c == 3:
                cell.alignment = Alignment(horizontal='center')
        if st != 'ביטל/ה':
            fill = GROUP_FILL if 'קבוצה' in (x['type'] or '') else (
                NEW_FILL if st == 'נרשם/ה אחרי המסרון' else None)
            if fill:
                for c in range(1, len(COLS) + 1):
                    ws.cell(row=i, column=c).fill = fill

    last = 5 + len(rows)
    ws.auto_filter.ref = 'A5:%s%d' % (get_column_letter(len(COLS)), last)
    ws.cell(row=last + 2, column=1, value='כל הטלפונים להעתקה (ללא מי שביטל):').font = \
        Font(bold=True, size=10, name='Assistant')
    ws.cell(row=last + 3, column=1,
            value=', '.join(x['phone'] for x in rows if status_of(x) != 'ביטל/ה')).font = \
        Font(size=10, name='Assistant')
    ws.cell(row=last + 5, column=1, value='רקע כתום = קבוצה או ארגון · רקע תכלת = נרשם אחרי שהמסרון יצא').font = SUB

    fname = '%02d - %s (%d משתתפים).xlsx' % (rank, safe_name(site), s['ppl'])
    wbx.save(os.path.join(OUTDIR, fname))
    written.append((fname, site, s['ppl'], len(rows), never, newreg, declined))

# ---- summary file ----
wbs = Workbook()
ws = wbs.active
ws.title = 'סיכום'
ws.sheet_view.rightToLeft = True
ws['A1'] = 'מוקדים ללא אף אישור הגעה'
ws['A1'].font = TITLE
ws['A2'] = 'קובץ נפרד לכל מוקד · נכון לייצוא 12/10/2026'
ws['A2'].font = SUB
for c, (h, w) in enumerate([('#', 5), ('מוקד', 40), ('משתתפים', 11), ('הרשמות', 10),
                            ('לא השיבו', 11), ('נרשמו אחרי המסרון', 18),
                            ('ביטלו', 9), ('שם הקובץ', 52)], start=1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    cell.alignment = Alignment(horizontal='center')
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = 'A5'
for i, (fname, site, ppl, nrows, never, newreg, declined) in enumerate(written, start=5):
    for c, v in enumerate([i - 4, site, ppl, nrows, never, newreg, declined, fname], start=1):
        cell = ws.cell(row=i, column=c, value=v)
        cell.font, cell.border = BODY, THIN
        if c not in (2, 8):
            cell.alignment = Alignment(horizontal='center')
wbs.save(os.path.join(OUTDIR, '00 - סיכום.xlsx'))

print('נכתבו %d קבצים לתיקייה %s\\' % (len(written) + 1, OUTDIR))
print('%-42s %5s %5s %6s %6s' % ('מוקד', 'משת׳', 'הרש׳', 'לא ענו', 'חדשים'))
print('-' * 70)
for fname, site, ppl, nrows, never, newreg, declined in written:
    print('%-42s %5d %5d %6d %6d' % (site[:42], ppl, nrows, never, newreg))
print('-' * 70)
print('%-42s %5d %5d %6d %6d' % ('סה״כ', sum(w[2] for w in written), sum(w[3] for w in written),
                                 sum(w[4] for w in written), sum(w[5] for w in written)))
